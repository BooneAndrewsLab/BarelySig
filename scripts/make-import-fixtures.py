#!/usr/bin/env python3
"""Writes the spreadsheet fixtures for opening a data file (item 10).

Each file is made by a writer independent of SheetJS, the reader under
test: openpyxl (.xlsx), odfpy (.ods) and xlwt (.xls). Run with any Python
that has them:

    pip install openpyxl odfpy xlwt
    python scripts/make-import-fixtures.py

Output: src/io/import/fixtures/*. The test (workbook.test.ts) states what
each file must read as, so change both together.
"""

import datetime
import pathlib

OUT = pathlib.Path(__file__).resolve().parent.parent / "src" / "io" / "import" / "fixtures"


def xlsx() -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Viability"
    ws.append(["WT", "KO", "Het"])
    ws.append([0.1 + 0.2, 2.5, 7])
    ws.append([1.23456789012345, "NA", -3])
    ws.append([1e-20, "#DIV/0!", True])
    ws.append([None, datetime.date(2026, 9, 26), "bad well"])
    ws.append([12, None, None])
    # Shown rounded; the stored value must come through in full.
    for row in ws.iter_rows(min_row=2, max_col=1):
        row[0].number_format = "0.00"

    pct = wb.create_sheet("Percent")
    pct.append(["Control", "Drug"])
    pct.append([0.85, 0.125])
    pct.append([0.9, 0.5])
    for row in pct.iter_rows(min_row=2):
        row[0].number_format = "0%"
        row[1].number_format = "0.0%"

    grouped = wb.create_sheet("Grouped")
    grouped.append([None, "WT", None, None, "KO", None, None])
    grouped.append([None, 1, 2, 3, 1, 2, 3])
    grouped.append(["Control", 1.1, 1.2, 1.3, 2.1, 2.2, 2.3])
    grouped.append(["Drug", 3.1, 3.2, None, 4.1, 4.2, 4.3])
    grouped.merge_cells("B1:D1")
    grouped.merge_cells("E1:G1")
    # A merged number is not copied into the cells it covers.
    grouped.append(["Merged", 9, None, None, None, None, None])
    grouped.merge_cells("B5:C5")

    hidden = wb.create_sheet("Hidden")
    hidden.append(["x"])
    hidden.append([1])
    hidden.sheet_state = "hidden"

    wb.create_sheet("Empty")
    wb.save(OUT / "workbook.xlsx")


def ods() -> None:
    from odf.opendocument import OpenDocumentSpreadsheet
    from odf.style import Style, TableColumnProperties
    from odf.number import PercentageStyle, Number, Text as NumberText
    from odf.table import Table, TableRow, TableCell, CoveredTableCell
    from odf.text import P

    doc = OpenDocumentSpreadsheet()
    pstyle = PercentageStyle(name="pct")
    pstyle.addElement(Number(decimalplaces="0", minintegerdigits="1"))
    pstyle.addElement(NumberText(text="%"))
    doc.automaticstyles.addElement(pstyle)
    pcell = Style(name="pctcell", family="table-cell", datastylename="pct")
    doc.automaticstyles.addElement(pcell)
    col = Style(name="co1", family="table-column")
    col.addElement(TableColumnProperties(columnwidth="2cm"))
    doc.automaticstyles.addElement(col)

    def text(s: str, span: int = 1) -> TableCell:
        c = TableCell(valuetype="string", numbercolumnsspanned=span) if span > 1 else TableCell(valuetype="string")
        c.addElement(P(text=s))
        return c

    def num(v: float) -> TableCell:
        c = TableCell(valuetype="float", value=v)
        c.addElement(P(text=str(v)))
        return c

    def pct(v: float) -> TableCell:
        c = TableCell(valuetype="percentage", value=v, stylename=pcell)
        c.addElement(P(text=f"{v * 100:g}%"))
        return c

    t = Table(name="Summary")
    rows = [
        [text("Group"), text("Mean"), text("SD"), text("N")],
        [text("WT"), num(5.5), num(1.25), num(4)],
        [text("KO"), num(3.25), num(0.5), num(5)],
    ]
    for cells in rows:
        tr = TableRow()
        for c in cells:
            tr.addElement(c)
        t.addElement(tr)
    doc.spreadsheet.addElement(t)

    g = Table(name="Merged")
    tr = TableRow()
    tr.addElement(TableCell())
    tr.addElement(text("WT", 2))
    tr.addElement(CoveredTableCell())
    tr.addElement(text("KO", 2))
    tr.addElement(CoveredTableCell())
    g.addElement(tr)
    for label, vals in [("A", [1.5, 2.5, 3.5, 4.5]), ("B", [5.5, 6.5, 7.5, 8.5])]:
        tr = TableRow()
        tr.addElement(text(label))
        for v in vals:
            tr.addElement(num(v))
        g.addElement(tr)
    tr = TableRow()
    tr.addElement(text("pct"))
    tr.addElement(pct(0.85))
    g.addElement(tr)
    doc.spreadsheet.addElement(g)
    doc.save(str(OUT / "workbook.ods"))


def xls() -> None:
    import xlwt

    wb = xlwt.Workbook(encoding="utf-8")
    ws = wb.add_sheet("Legacy")
    ws.write_merge(0, 0, 1, 2, "Wild type")
    ws.write_merge(0, 0, 3, 4, "Mutant ü")
    for j, v in enumerate([1, 2, 1, 2]):
        ws.write(1, j + 1, v)
    ws.write(2, 0, "Day 1")
    for j, v in enumerate([0.5, 0.25, 1.75, 2.125]):
        ws.write(2, j + 1, v)
    ws.write(3, 0, "Day 2")
    for j, v in enumerate([10, 20, 30, 40]):
        ws.write(3, j + 1, v)
    wb.save(str(OUT / "legacy.xls"))


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    xlsx()
    ods()
    xls()
    print(f"wrote fixtures to {OUT}")
